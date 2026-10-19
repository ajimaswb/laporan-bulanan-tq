#!/usr/bin/env python3
"""Generator Laporan Bulanan FM-SDD-008 dari spreadsheet kehadiran.

Contoh:
    python3 laporan_bulanan.py                  # periode berjalan (cutoff tgl 25)
    python3 laporan_bulanan.py --periode 2026-10  # 26 Sep - 25 Okt 2026
    python3 laporan_bulanan.py --sumber kehadiran.xlsx  # pakai file lokal, tanpa download
    python3 laporan_bulanan.py --sdo JNT         # laporan untuk SDO lain
"""
import argparse
import datetime as dt
import json
import re
from copy import copy
import subprocess
import sys
import tempfile
import urllib.request
from io import BytesIO
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image as XlImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor, TwoCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import PatternFill
from PIL import Image as PILImage

import penyimpanan

BASE = Path(__file__).resolve().parent

BULAN = ["JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI",
         "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER"]
HARI_ID = ["SENIN", "SELASA", "RABU", "KAMIS", "JUMAT", "SABTU", "MINGGU"]
HARI_EN = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

# Posisi di template
BARIS_DATA = 8          # baris pertama tabel kegiatan
MAKS_BARIS = 23         # jumlah baris data di template (8..30)
KOLOM = {"no": "A", "hari": "C", "tanggal": "F", "copart": "K", "kegiatan": "O", "ket": "AD"}
BLOK_MANDAYS = [11, 37, 63, 89]  # baris "Nama Copart n" di sheet Laporan Mandays
SEL_DIBUAT = "A31"      # "Dibuat,\n\n\n<nama SDO>" di sheet Laporan Bulanan
TTD_MAKS = (1_300_000, 760_000)  # ukuran maks tanda tangan pembuat (EMU)

KUNING = PatternFill("solid", fgColor="FFFF00")
AKRONIM = {"QA", "ISO", "AMI", "EMS", "QMS", "RTM", "TQ", "SPI", "DMOS", "IMS", "K3"}


# ---------------------------------------------------------------- periode
def hitung_periode(periode: str | None, cutoff: int, hari_ini: dt.date):
    """Periode diberi nama bulan akhirnya: 2026-10 = 26 Sep 2026 s/d 25 Okt 2026."""
    if periode:
        th, bl = map(int, periode.split("-"))
    else:
        th, bl = hari_ini.year, hari_ini.month
        if hari_ini.day > cutoff:
            th, bl = (th + 1, 1) if bl == 12 else (th, bl + 1)
    akhir = dt.date(th, bl, cutoff)
    awal_bl, awal_th = (12, th - 1) if bl == 1 else (bl - 1, th)
    awal = dt.date(awal_th, awal_bl, cutoff) + dt.timedelta(days=1)
    return awal, akhir


def label_periode(awal: dt.date, akhir: dt.date) -> str:
    a = f"{awal.day} {BULAN[awal.month - 1].title()}"
    if awal.year != akhir.year:
        a += f" {awal.year}"
    return f"{a} - {akhir.day} {BULAN[akhir.month - 1].title()} {akhir.year}"


# ---------------------------------------------------------------- sumber data
def unduh_spreadsheet(sheet_id: str, tujuan: Path):
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    try:
        with urllib.request.urlopen(url, timeout=60, context=penyimpanan.SSL) as r:
            tujuan.write_bytes(r.read())
    except Exception:
        # Python bawaan macOS kadang belum punya sertifikat SSL; pakai curl.
        subprocess.run(["curl", "-sSLf", "-o", str(tujuan), url], check=True)


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().upper()


class Kehadiran:
    """Membaca sheet bulanan (HARI | TANGGAL | kode implementer ...)."""

    def __init__(self, path: Path):
        self.wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        self._cache = {}

    def _cari_sheet(self, tahun: int, bulan: int):
        nama = BULAN[bulan - 1]
        kandidat = []
        for ws in self.wb.worksheets:
            t = norm(ws.title)
            if t == f"{nama} {tahun}":
                return ws
            if t == nama:
                kandidat.append(ws)
        # Sheet tanpa tahun (mis. "OKTOBER"): pastikan nama hari tgl 1 cocok.
        hari1 = HARI_ID[dt.date(tahun, bulan, 1).weekday()]
        for ws in kandidat:
            for row in ws.iter_rows(min_row=2, max_row=6, max_col=2, values_only=True):
                if row[1] in (1, 1.0):
                    if norm(row[0]) == hari1:
                        return ws
                    break
        return None

    def bulan(self, tahun: int, bulan: int):
        """-> {tanggal: {kode: teks}} untuk satu bulan."""
        key = (tahun, bulan)
        if key in self._cache:
            return self._cache[key]
        ws = self._cari_sheet(tahun, bulan)
        data = {}
        if ws is not None:
            rows = list(ws.iter_rows(values_only=True))
            header = [norm(h) for h in rows[0]]
            for row in rows[1:]:
                if len(row) < 2 or not isinstance(row[1], (int, float)):
                    continue
                try:
                    tgl = dt.date(tahun, bulan, int(row[1]))
                except ValueError:
                    continue
                data[tgl] = {header[i]: norm(v) for i, v in enumerate(row)
                             if i >= 2 and i < len(header) and header[i] and v not in (None, "")}
        self._cache[key] = data
        return data

    def implementer(self, tahun, bulan):
        """Kode implementer di header sheet bulan tsb."""
        ws = self._cari_sheet(tahun, bulan)
        if ws is None:
            return []
        header = next(ws.iter_rows(max_row=1, values_only=True))
        return [norm(h) for h in header[2:] if norm(h)]

    def ada_sheet(self, tahun, bulan):
        return self._cari_sheet(tahun, bulan) is not None

    def hari(self, tgl: dt.date):
        return self.bulan(tgl.year, tgl.month).get(tgl, {})


# ---------------------------------------------------------------- klasifikasi
class Klasifikasi:
    def __init__(self, cfg):
        self.cfg = cfg
        self.alias = {}  # kata kunci (boleh lebih dari 1 kata) -> kode
        for kode, c in cfg["copart"].items():
            for k in [kode, *c.get("kata_kunci", [])]:
                k = norm(k)
                if k:
                    self.alias[k] = kode
        alt = "|".join(re.escape(a) for a in sorted(self.alias, key=len, reverse=True))
        self.pola = re.compile(rf"(?<![A-Z0-9])(?:{alt})(?![A-Z0-9])") if alt else None

    def cari_copart(self, teks: str):
        """Kode copart yang paling awal disebut di teks."""
        m = self.pola.search(teks) if self.pola else None
        return self.alias[m.group(0)] if m else None

    def buang_alias(self, teks: str) -> str:
        return self.pola.sub(" ", teks) if self.pola else teks

    def angka_mandays(self, teks: str, kode: str):
        if self.cari_copart(teks) != kode:
            return None
        angka = re.findall(r"\b(\d+)\b", teks)
        return int(angka[-1]) if angka else None

    def kegiatan(self, teks: str):
        """-> dict(copart, kegiatan, ket, kategori, ragu)"""
        if not teks:
            return dict(copart="?", kegiatan="(kosong di spreadsheet)", ket="?", kategori="?", ragu=True)
        for kunci, label in self.cfg["absen"].items():
            if re.search(rf"\b{kunci}\b", teks):
                return dict(copart="-", kegiatan=label, ket="-", kategori="absen", ragu=False)

        kode = self.cari_copart(teks)
        if not kode:
            return dict(copart="?", kegiatan=judul(re.sub(r"\b\d+\b", "", teks)), ket="?", kategori="?", ragu=True)

        c = self.cfg["copart"][kode]
        kata = [w for w in re.findall(r"[A-Z0-9]+", teks) if not w.isdigit() and w != "DAN"]
        sisa = [w for w in re.findall(r"[A-Z0-9]+", self.buang_alias(teks)) if not w.isdigit() and w != "DAN"]
        penuh, frasa = " ".join(kata), " ".join(sisa)
        khusus = {k.upper(): v for k, v in c.get("kegiatan_khusus", {}).items()}
        std = c.get("standar", "")
        ragu = not c.get("lokasi")

        if penuh in khusus:
            keg = khusus[penuh]
        elif frasa in khusus:
            keg = khusus[frasa]
        elif not sisa:
            keg = c.get("kegiatan_default") or f"Pendampingan {std} {kode}"
            ragu = ragu or c.get("cek_jika_default", False)
        elif "TRAINING" in sisa:
            keg = f"Training AMI {kode}" if "AMI" in sisa else f"Training {std} {kode}"
        elif "QA" in sisa and "CLOSING" in sisa:
            keg = f"QA & Closing {kode}"
        elif "QA" in sisa:
            keg = f"QA {kode}"
        elif "CLOSING" in sisa:
            keg = f"Closing {kode}"
        elif "RTM" in sisa:
            keg = f"RTM {kode}"
        elif sisa[:2] == ["AUDIT", "KEUANGAN"]:
            keg = f"Audit Keuangan {kode}"
        elif sisa[0] in ("MAINTENANCE", "MTC"):
            keg = f"Maintenance {kode}"
        else:
            keg = f"{judul(frasa)} {kode}"

        return dict(copart=kode, kegiatan=re.sub(r"\s+", " ", keg).strip(), ket=c.get("lokasi") or "?",
                    kategori=c.get("kategori", "copart"), ragu=ragu)

    def koreksi(self, b: dict, isi: dict):
        """Terapkan koreksi manual dari config untuk satu tanggal."""
        b.update({k: v for k, v in isi.items() if k in ("copart", "kegiatan", "ket")})
        if "copart" in isi:
            c = self.cfg["copart"].get(isi["copart"], {})
            b["kategori"] = c.get("kategori", "copart") if isi["copart"] != "-" else "absen"
            b.setdefault("ket", c.get("lokasi") or "?")
        b["ragu"] = b["copart"] == "?" or b["ket"] == "?"


def judul(s: str) -> str:
    out = []
    for w in s.split():
        if w in AKRONIM:
            out.append(w)
        elif w == "DAN":
            out.append("dan")
        else:
            out.append(w.capitalize())
    return " ".join(out)


# ---------------------------------------------------------------- penulisan xlsx
def hapus_baris(ws, idx: int, jumlah: int):
    """delete_rows yang juga menggeser merge, tinggi baris, dan gambar (tanda tangan)."""
    if jumlah <= 0:
        return
    batas = idx + jumlah
    merges = list(ws.merged_cells.ranges)
    for m in merges:
        ws.merged_cells.remove(m)
    tinggi = {r: ws.row_dimensions[r].height for r in range(batas, ws.max_row + 1)}
    emu_dihapus = sum(int((ws.row_dimensions[r].height or 15) * 12700) for r in range(idx, batas))
    ws.delete_rows(idx, jumlah)
    for m in merges:
        if m.max_row < idx:
            ws.merged_cells.add(m)
        elif m.min_row >= batas:
            m.shift(row_shift=-jumlah)
            ws.merged_cells.add(m)
    for r, h in tinggi.items():
        ws.row_dimensions[r - jumlah].height = h
    for img in ws._images:
        a = img.anchor
        if a._from.row < batas - 1:
            continue
        for titik in (a._from, getattr(a, "to", None)):
            if titik is not None:
                titik.row -= jumlah
        # Posisi absolut (dipakai preview macOS/Google Sheets) ikut digeser
        xfrm = a.pic.spPr.xfrm if getattr(a, "pic", None) and a.pic.spPr else None
        if xfrm is not None and xfrm.off is not None:
            xfrm.off.y = max(0, xfrm.off.y - emu_dihapus)


def pasang_ttd(ws, path):
    """Ganti tanda tangan pembuat (gambar kiri-bawah) dengan milik SDO; kosongkan jika tidak ada."""
    lama = [im for im in ws._images
            if isinstance(im.anchor, TwoCellAnchor) and im.anchor._from.col <= 8 and im.anchor._from.row >= 20]
    for im in lama:
        ws._images.remove(im)
    if not lama or not path or not Path(path).exists():
        return
    asal = lama[0].anchor._from
    gambar = PILImage.open(path).convert("RGBA")
    kotak = gambar.getchannel("A").getbbox()
    if kotak:
        gambar = gambar.crop(kotak)
    skala = min(TTD_MAKS[0] / gambar.width, TTD_MAKS[1] / gambar.height)
    cx, cy = int(gambar.width * skala), int(gambar.height * skala)
    buf = BytesIO()
    gambar.save(buf, format="PNG")
    img = XlImage(BytesIO(buf.getvalue()))
    geser = 120_000
    anchor = OneCellAnchor(_from=AnchorMarker(col=asal.col, colOff=asal.colOff, row=asal.row, rowOff=asal.rowOff + geser),
                           ext=XDRPositiveSize2D(cx, cy))
    # Pakai bingkai gambar lama agar posisi absolut (xfrm) ikut terisi untuk preview macOS/Google Sheets
    pic = lama[0].anchor.pic
    if pic is not None and pic.spPr is not None and pic.spPr.xfrm is not None:
        pic.spPr.xfrm.ext.cx, pic.spPr.xfrm.ext.cy = cx, cy
        pic.spPr.xfrm.off.y += geser
        anchor.pic = pic
    ws.add_image(img, anchor)


def tulis_laporan(cfg, awal, akhir, baris, ringkas, mandays, out_path: Path):
    wb = openpyxl.load_workbook(BASE / cfg["template"])
    bulan_label = f"{BULAN[akhir.month - 1].title()} {akhir.year}"
    kode_imp = cfg["kode_implementer"]

    # --- Sheet 1: Laporan Bulanan
    ws = wb["Laporan Bulanan"]
    ws["I5"] = kode_imp
    ws["Y5"] = bulan_label
    ws[SEL_DIBUAT] = f"Dibuat,\n\n\n{cfg.get('nama_sdo') or ''}"
    pasang_ttd(ws, cfg.get("ttd"))
    if len(baris) > MAKS_BARIS:
        raise ValueError(f"Jumlah hari kerja {len(baris)} > {MAKS_BARIS} baris template.")
    for i, b in enumerate(baris):
        r = BARIS_DATA + i
        ws[f"{KOLOM['no']}{r}"] = i + 1
        sel = ws[f"{KOLOM['hari']}{r}"]
        sel.value = HARI_EN[b["tanggal"].weekday()]
        font = copy(sel.font)
        font.b = b["tanggal"].weekday() == 0  # hanya Senin yang tebal
        sel.font = font
        ws[f"{KOLOM['tanggal']}{r}"] = dt.datetime.combine(b["tanggal"], dt.time())
        ws[f"{KOLOM['copart']}{r}"] = b["copart"]
        ws[f"{KOLOM['kegiatan']}{r}"] = b["kegiatan"]
        ws[f"{KOLOM['ket']}{r}"] = b["ket"]
        if b["ragu"]:
            for k in ("copart", "kegiatan", "ket"):
                ws[f"{KOLOM[k]}{r}"].fill = KUNING
    hapus_baris(ws, BARIS_DATA + len(baris), MAKS_BARIS - len(baris))
    baris_akhir = BARIS_DATA + len(baris) - 1

    # --- Sheet 2: Laporan Mandays
    ws = wb["Laporan Mandays"]
    ws["I8"] = kode_imp
    ws["Y8"] = bulan_label
    pasang_ttd(ws, cfg.get("ttd"))
    for n, r in enumerate(BLOK_MANDAYS):
        m = mandays[n] if n < len(mandays) else None
        ws[f"G{r}"] = m["kode"] if m else None
        ws[f"AA{r}"] = m["terpakai"] if m else "N/A"
        ws[f"I{r + 2}"] = m["progress_step"] if m else None
        ws[f"AA{r + 2}"] = f"={m['kuota']}-AA{r}" if m else "N/A"
        ws[f"B{r + 6}"] = m["laporan_singkat"] if m else None
    rng = f"'Laporan Bulanan'!A{BARIS_DATA}:B{baris_akhir}"
    for r, key in ((114, "copart"), (116, "lain"), (118, "kantor")):
        ws[f"S{r}"] = ringkas[key]
        ws[f"V{r}"] = f"=COUNT({rng})"

    wb.calculation.fullCalcOnLoad = True
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


# ---------------------------------------------------------------- penyusunan
def _store(store=None):
    return store or penyimpanan.PenyimpananLokal(BASE)


def baca_config(path=None, sdo=None, store=None) -> dict:
    """Config global + master copart + data milik satu SDO (dari file lokal atau Supabase)."""
    st = _store(store)
    cfg = json.loads(Path(path or BASE / "config.json").read_text())
    master = st.baca_master()
    cfg["copart"] = master["copart"]
    cfg["abaikan"] = master["abaikan"]
    kode = norm(sdo or getattr(st, "kode", None) or cfg.get("sdo_default") or "AJM")
    data = st.baca_sdo().get(kode, {})
    cfg["kode_implementer"] = kode
    cfg["nama_sdo"] = data.get("nama", "")
    cfg["mandays_copart"] = data.get("mandays_copart", [])
    cfg["koreksi"] = data.get("koreksi", {})
    cfg["ttd"] = st.ttd_path(kode)
    return cfg


def simpan_config(cfg: dict, path=None, store=None):
    """Kebalikan baca_config: folder output ke config.json, bagian SDO ke penyimpanan."""
    st = _store(store)
    if isinstance(st, penyimpanan.PenyimpananLokal):
        path = Path(path or BASE / "config.json")
        global_ = json.loads(path.read_text())
        if global_.get("folder_output") != cfg.get("folder_output"):
            global_["folder_output"] = cfg["folder_output"]
            path.write_text(json.dumps(global_, indent=2, ensure_ascii=False) + "\n")
    st.simpan_sdo(cfg["kode_implementer"], {"nama": cfg.get("nama_sdo", ""),
                                             "mandays_copart": cfg.get("mandays_copart", []),
                                             "koreksi": cfg.get("koreksi", {})})


def belum_ada_lokasi(cfg) -> list:
    return sorted(k for k, c in cfg["copart"].items() if not (c.get("lokasi") or "").strip())


def deteksi_copart_baru(cfg, keh: "Kehadiran", awal, akhir) -> list:
    """Nama di spreadsheet (semua implementer, bulan-bulan periode ini) yang belum ada di master."""
    det = cfg.get("deteksi", {})
    kota = {norm(k): v for k, v in det.get("kota", {}).items()}
    stop = {norm(k) for k in det.get("kata_kegiatan", [])} | {"DAN"}
    absen = [norm(k) for k in cfg["absen"]]
    abaikan = {norm(a) for a in cfg.get("abaikan", [])}
    kls = Klasifikasi(cfg)

    bulan = sorted({(awal.year, awal.month), (akhir.year, akhir.month)})
    temuan = {}
    for th, bl in bulan:
        data = keh.bulan(th, bl)
        implementer = {k for isi in data.values() for k in isi}
        for tgl, isi in sorted(data.items()):
            for imp, teks in isi.items():
                for seg in re.split(r"[,/+&()]|-", teks):
                    if any(re.search(rf"\b{a}\b", seg) for a in absen):
                        continue
                    kata = re.findall(r"[A-Z][A-Z0-9]*", kls.buang_alias(norm(seg)))
                    lokasi = next((kota[w] for w in kata if w in kota), "")
                    sisa = [w for w in kata if w not in stop and w not in kota and w not in implementer]
                    frasa = " ".join(sisa)
                    if len(frasa) < 2 or frasa in abaikan:
                        continue
                    t = temuan.setdefault(frasa, {"frasa": frasa, "kode": sisa[0], "lokasi": lokasi,
                                                  "jumlah": 0, "contoh": teks, "implementer": set(),
                                                  "pertama": tgl.isoformat()})
                    t["jumlah"] += 1
                    t["implementer"].add(imp)
                    t["lokasi"] = t["lokasi"] or lokasi
    hasil = sorted(temuan.values(), key=lambda t: (-t["jumlah"], t["frasa"]))
    for t in hasil:
        t["implementer"] = sorted(t["implementer"])
    return hasil


def nama_file(cfg, awal, akhir) -> str:
    return f"FM-SDD-008 Laporan Bulanan {cfg['kode_implementer']} ({label_periode(awal, akhir)}).xlsx"


def folder_output(cfg) -> Path:
    p = Path(cfg["folder_output"]).expanduser()
    return p if p.is_absolute() else BASE / p


def hitung_ringkas(baris, cfg) -> dict:
    utama = {m["kode"] for m in cfg["mandays_copart"]}
    ringkas = {"copart": 0, "lain": 0, "kantor": 0, "absen": 0}
    for b in baris:
        if b["kategori"] == "absen":
            ringkas["absen"] += 1
        elif b["kategori"] == "kantor":
            ringkas["kantor"] += 1
        elif b["copart"] in utama:
            ringkas["copart"] += 1
        elif b["kategori"] != "?":
            ringkas["lain"] += 1
    return ringkas


def susun_laporan(cfg, keh: Kehadiran, awal, akhir) -> dict:
    """Hitung semua isi laporan. Tiap baris menyimpan hasil otomatis di 'auto'."""
    kode_imp = cfg["kode_implementer"].upper()
    kls = Klasifikasi(cfg)
    peringatan = [f"Sheet {BULAN[bl - 1]} {th} tidak ditemukan di spreadsheet"
                  for th, bl in sorted({(awal.year, awal.month), (akhir.year, akhir.month)})
                  if not keh.ada_sheet(th, bl)]

    # Baris laporan: semua hari kerja (Senin-Jumat) dalam periode
    baris = []
    tgl = awal
    while tgl <= akhir:
        teks = keh.hari(tgl).get(kode_imp, "")
        if tgl.weekday() < 5:
            b = kls.kegiatan(teks)
            b["auto"] = {k: b[k] for k in ("copart", "kegiatan", "ket")}
            b["dikoreksi"] = tgl.isoformat() in cfg.get("koreksi", {})
            if b["dikoreksi"]:
                kls.koreksi(b, cfg["koreksi"][tgl.isoformat()])
            b["tanggal"], b["teks"] = tgl, teks
            baris.append(b)
            if b["ragu"]:
                peringatan.append(f"{tgl:%d-%m-%Y}: '{teks}' perlu dicek")
        elif teks:
            peringatan.append(f"{tgl:%d-%m-%Y} (akhir pekan, tidak masuk laporan): '{teks}'")
        tgl += dt.timedelta(days=1)

    # Mandays terpakai: angka terakhir yang tercatat utk copart tsb (semua implementer)
    mandays = []
    for m in cfg["mandays_copart"]:
        terpakai = m.get("mandays_terpakai")
        if terpakai is None:
            terpakai, kapan = 0, None
            th, bl = akhir.year, akhir.month
            for _ in range(12):  # telusuri mundur maks 12 bulan
                for d, isi in sorted(keh.bulan(th, bl).items(), reverse=True):
                    if d > akhir:
                        continue
                    nilai = [x for x in (kls.angka_mandays(t, m["kode"]) for t in isi.values()) if x]
                    if nilai:
                        terpakai, kapan = max(nilai), d
                        break
                if kapan:
                    break
                th, bl = (th - 1, 12) if bl == 1 else (th, bl - 1)
            info = f"angka terakhir {kapan:%d-%m-%Y}" if kapan else "tidak ditemukan"
        else:
            info = "diisi manual"
        mandays.append({**m, "terpakai": terpakai, "info": info})

    return dict(awal=awal, akhir=akhir, baris=baris, ringkas=hitung_ringkas(baris, cfg),
                mandays=mandays, peringatan=peringatan)


def muat_kehadiran(cfg, sumber=None, folder_tmp=None) -> Kehadiran:
    if sumber:
        return Kehadiran(Path(sumber))
    src = Path(folder_tmp or tempfile.mkdtemp()) / "kehadiran.xlsx"
    unduh_spreadsheet(cfg["spreadsheet_id"], src)
    return Kehadiran(src)


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--periode", help="YYYY-MM bulan akhir periode (default: periode berjalan)")
    ap.add_argument("--sumber", help="file xlsx kehadiran lokal (default: unduh dari Google Sheets)")
    ap.add_argument("--sdo", help="kode SDO (default: sdo_default di config)")
    ap.add_argument("--config", default=str(BASE / "config.json"))
    ap.add_argument("--output", help="folder output (default: dari config)")
    args = ap.parse_args()

    cfg = baca_config(args.config, args.sdo)
    print(f"SDO: {cfg['kode_implementer']} {cfg['nama_sdo']}")
    awal, akhir = hitung_periode(args.periode, cfg["tanggal_cutoff"], dt.date.today())
    print(f"Periode: {label_periode(awal, akhir)}")
    with tempfile.TemporaryDirectory() as tmp:
        if not args.sumber:
            print("Mengunduh spreadsheet kehadiran...")
        keh = muat_kehadiran(cfg, args.sumber, tmp)
        lap = susun_laporan(cfg, keh, awal, akhir)

    out_dir = Path(args.output) if args.output else folder_output(cfg)
    out_path = out_dir / nama_file(cfg, awal, akhir)
    tulis_laporan(cfg, awal, akhir, lap["baris"], lap["ringkas"], lap["mandays"], out_path)

    for m in lap["mandays"]:
        print(f"  Mandays {m['kode']}: {m['terpakai']}/{m['kuota']} ({m['info']})")
    print(f"\n{'No':>3} {'Tanggal':<11} {'Co-Part':<7} {'Kegiatan':<32} {'Keterangan':<12} | Spreadsheet")
    for i, b in enumerate(lap["baris"], 1):
        tanda = " <- CEK" if b["ragu"] else ""
        print(f"{i:>3} {b['tanggal']:%d-%m-%Y}  {b['copart']:<7} {b['kegiatan']:<32} {b['ket']:<12} | {b['teks']}{tanda}")
    r = lap["ringkas"]
    print(f"\nMandays: copart={r['copart']} lain/non-mandays={r['lain']} "
          f"kantor={r['kantor']} absen={r['absen']} / {len(lap['baris'])} hari kerja")
    if lap["peringatan"]:
        print("\nPerlu dicek:")
        for p in lap["peringatan"]:
            print("  -", p)
    print(f"\nTersimpan: {out_path}")


if __name__ == "__main__":
    main()
